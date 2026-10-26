"""
Shared helpers for the export endpoints: request validation, XLSX and PDF building.
"""

import io
from datetime import datetime

from django.http import HttpResponse
from rest_framework import status
from rest_framework.negotiation import DefaultContentNegotiation
from rest_framework.response import Response

from apps.companies.models import Company
from apps.users.company_access import get_accessible_company_ids


# Optional heavy-import guard — openpyxl / reportlab are added to
# requirements.txt; these guards give a readable error if somehow missing.
# ---------------------------------------------------------------------------
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    OPENPYXL_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    REPORTLAB_AVAILABLE = True
except ImportError:  # pragma: no cover
    REPORTLAB_AVAILABLE = False


# ============================================================
# SHARED HELPERS
# ============================================================

def _bad_request(msg: str):
    return Response({"error": msg}, status=status.HTTP_400_BAD_REQUEST)


def _forbidden(msg: str):
    return Response({"detail": msg}, status=status.HTTP_403_FORBIDDEN)


def _resolve_companies(company_id_strs, user):
  
    if not company_id_strs:
        accessible_ids = get_accessible_company_ids(user)
        qs = Company.objects.filter(is_active=True)
        if accessible_ids is not None:
            qs = qs.filter(pk__in=accessible_ids)
        return list(qs.order_by("symbol")), None
    company_ids = []
    for s in company_id_strs:
        if not s:
            continue
        try:
            company_ids.append(int(s))
        except (ValueError, TypeError):
            return None, _bad_request("company_id must be an integer.")
    try:
        companies = list(
            Company.objects.filter(pk__in=company_ids, is_active=True).order_by("symbol")
        )
    except Exception:
        return None, _bad_request("Invalid company IDs.")
    found_ids = {c.id for c in companies}
    accessible_ids = get_accessible_company_ids(user)
    if accessible_ids is not None:
        denied = [cid for cid in company_ids if cid not in set(accessible_ids)]
        if denied:
            return None, _forbidden("You do not have access to one or more requested companies.")
    for cid in company_ids:
        if cid not in found_ids:
            return None, _bad_request(
                f"Company with id={cid} does not exist or is not active."
            )
    return companies, None


def _parse_date(value: str, field_name: str):
    """Parse YYYY-MM-DD string. Returns (date_obj, error_response)."""
    if not value:
        return None, None
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date(), None
    except ValueError:
        return None, _bad_request(
            f"Invalid {field_name} format. Use YYYY-MM-DD."
        )


def _validate_date_range(date_from, date_to):
    """Returns error response if date_from > date_to, else None."""
    if date_from and date_to and date_from > date_to:
        return _bad_request("date_from must not be after date_to.")
    return None


def _filename(prefix, companies, date_from, date_to, ext):
    """Build a clean download filename."""
    parts = [prefix]
    if companies:
        parts.append("_".join(c.symbol for c in companies))
    if date_from:
        parts.append(str(date_from))
    if date_to:
        parts.append(f"to_{date_to}")
    return "_".join(parts) + "." + ext


# ============================================================
# XLSX HELPERS
# ============================================================

def _xlsx_header_style():
    """Bold white text on dark blue background."""
    fill = PatternFill("solid", fgColor="1F3864")
    font = Font(bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return fill, font, alignment


def _apply_xlsx_header(ws, headers: list[str], row: int = 1):
    fill, font, alignment = _xlsx_header_style()
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
    ws.row_dimensions[row].height = 18


def _autofit_xlsx(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _xlsx_response(wb, filename: str) -> HttpResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ============================================================
# PDF HELPERS
# ============================================================

def _pdf_doc(buf, title: str):
    """Return a SimpleDocTemplate bound to an in-memory buffer."""
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=title,
    )
    return doc


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        fontSize=16,
        fontName="Helvetica-Bold",
        spaceAfter=6,
        textColor=colors.HexColor("#1F3864"),
    ))
    styles.add(ParagraphStyle(
        name="SectionHeading",
        fontSize=12,
        fontName="Helvetica-Bold",
        spaceBefore=12,
        spaceAfter=4,
        textColor=colors.HexColor("#1F3864"),
    ))
    styles.add(ParagraphStyle(
        name="SubHeading",
        fontSize=10,
        fontName="Helvetica-Bold",
        spaceBefore=8,
        spaceAfter=2,
        textColor=colors.HexColor("#374151"),
    ))
    styles.add(ParagraphStyle(
        name="BodySmall",
        fontSize=9,
        fontName="Helvetica",
        spaceAfter=2,
    ))
    return styles


def _pdf_table_style(header_rows: int = 1):
    header_bg = colors.HexColor("#1F3864")
    alt_row_bg = colors.HexColor("#F3F4F6")
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), header_bg),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, header_rows - 1), 8),
        ("ALIGN", (0, 0), (-1, header_rows - 1), "CENTER"),
        ("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, alt_row_bg]),
        ("FONTNAME", (0, header_rows), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, header_rows), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])


def _pdf_response(buf, filename: str) -> HttpResponse:
    response = HttpResponse(buf.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class ExportContentNegotiation(DefaultContentNegotiation):
    """Treat the export ``format`` query parameter as file format, not renderer."""

    def select_renderer(self, request, renderers, format_suffix=None):
        django_request = request._request
        original_query = django_request.GET
        if "format" not in original_query:
            return super().select_renderer(request, renderers, format_suffix)

        django_request.GET = original_query.copy()
        django_request.GET.pop("format", None)
        try:
            return super().select_renderer(request, renderers, format_suffix)
        finally:
            django_request.GET = original_query
